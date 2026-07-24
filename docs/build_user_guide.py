"""Build the illustrated beginner handbook for the BMC Smart Log Platform.

This source is intentionally kept with the handbook so a future maintainer can
refresh the platform version, links, or screenshots without reformatting the
whole document manually.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "docs" / "assets" / "user-guide"
OUTPUT = ROOT / "deliverables" / "BMC智能日志平台_同事使用指导书_v0.6.0.docx"


# Palette: deliberately calm and high-contrast for printing and on-screen use.
NAVY = "17324D"
BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
CYAN = "1FA7C9"
PALE_BLUE = "EAF4FB"
PALE_CYAN = "E8F8FA"
PALE_GOLD = "FFF5DE"
PALE_RED = "FCEBEC"
PALE_GREEN = "EAF7F0"
WHITE = "FFFFFF"
INK = "213547"
MUTED = "5C6E7D"
LINE = "C9DCEB"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    """Set per-cell borders; kwargs use Word names such as top/bottom."""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, edge_data in kwargs.items():
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn("w:" + key), str(edge_data[key]))


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_column_width(cell, width_inches: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_inches)


def lock_table_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_east_asia_font(run, font_name="Heiti SC"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_east_asia_font(run)


def add_para(doc_or_cell, text="", size=10.5, color=INK, bold=False, align=None, before=0, after=6, line=1.25):
    paragraph = doc_or_cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    set_east_asia_font(run)
    return paragraph


def add_rich_para(doc_or_cell, parts, size=10.5, color=INK, after=6, before=0, line=1.25):
    paragraph = doc_or_cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    for text, is_bold, run_color in parts:
        run = paragraph.add_run(text)
        run.bold = is_bold
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(run_color or color)
        set_east_asia_font(run)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        paragraph.paragraph_format.space_before = Pt(18)
        paragraph.paragraph_format.space_after = Pt(10)
        size, color = 16, BLUE
    elif level == 2:
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(7)
        size, color = 13, BLUE
    else:
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)
        size, color = 12, DEEP_BLUE
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    set_east_asia_font(run)
    return paragraph


def add_rule(doc, color=LINE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lock_table_layout(table)
    cell = table.cell(0, 0)
    set_column_width(cell, 6.5)
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=12, bottom=12, start=0, end=0)
    cell.text = ""
    return table


def add_label(doc, text, color=CYAN):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(color)
    set_east_asia_font(run)
    return p


def add_callout(doc, title, body, tone="blue"):
    tones = {
        "blue": (PALE_BLUE, BLUE),
        "green": (PALE_GREEN, "2F8F68"),
        "gold": (PALE_GOLD, "B77A18"),
        "red": (PALE_RED, "B34D5D"),
    }
    fill, accent = tones[tone]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lock_table_layout(table)
    cell = table.cell(0, 0)
    set_column_width(cell, 6.5)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=180, end=180)
    set_cell_border(cell, left={"val": "single", "sz": 18, "color": accent})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = rgb(accent)
    set_east_asia_font(r)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    r2.font.size = Pt(9.7)
    r2.font.color.rgb = rgb(INK)
    set_east_asia_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_steps(doc, steps):
    for index, (title, detail) in enumerate(steps, 1):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        lock_table_layout(table)
        table.autofit = False
        number_cell, content_cell = table.rows[0].cells
        set_column_width(number_cell, 0.52)
        set_column_width(content_cell, 5.98)
        set_cell_shading(number_cell, BLUE)
        set_cell_shading(content_cell, WHITE)
        set_cell_margins(number_cell, top=95, bottom=95, start=60, end=60)
        set_cell_margins(content_cell, top=80, bottom=80, start=130, end=130)
        set_cell_border(content_cell, top={"val": "single", "sz": 8, "color": LINE}, bottom={"val": "single", "sz": 8, "color": LINE}, right={"val": "single", "sz": 8, "color": LINE})
        number_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p_no = number_cell.paragraphs[0]
        p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_no.paragraph_format.space_after = Pt(0)
        r_no = p_no.add_run(str(index))
        r_no.bold = True
        r_no.font.size = Pt(13)
        r_no.font.color.rgb = rgb(WHITE)
        set_east_asia_font(r_no)
        p_body = content_cell.paragraphs[0]
        p_body.paragraph_format.space_after = Pt(2)
        r_title = p_body.add_run(title)
        r_title.bold = True
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = rgb(DEEP_BLUE)
        set_east_asia_font(r_title)
        p_detail = content_cell.add_paragraph()
        p_detail.paragraph_format.space_after = Pt(0)
        p_detail.paragraph_format.line_spacing = 1.18
        r_detail = p_detail.add_run(detail)
        r_detail.font.size = Pt(9.6)
        r_detail.font.color.rgb = rgb(INK)
        set_east_asia_font(r_detail)
        prevent_row_split(table.rows[0])
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(item)
        run.font.size = Pt(10.2)
        run.font.color.rgb = rgb(INK)
        set_east_asia_font(run)


def add_image(doc, path: Path, width_inches: float, caption: str, alt_text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width_inches))
    inline._inline.docPr.set("descr", alt_text)
    inline._inline.docPr.set("title", alt_text)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    cr = cp.add_run(caption)
    cr.italic = True
    cr.font.size = Pt(8.7)
    cr.font.color.rgb = rgb(MUTED)
    set_east_asia_font(cr)


def add_module_cards(doc, cards):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lock_table_layout(table)
    table.autofit = False
    for row_index in range(0, len(cards), 2):
        row = table.add_row()
        for col_index in range(2):
            cell = row.cells[col_index]
            set_column_width(cell, 3.2)
            set_cell_margins(cell, top=125, bottom=125, start=150, end=130)
            set_cell_shading(cell, PALE_BLUE if (row_index + col_index) % 3 else PALE_CYAN)
            set_cell_border(cell, top={"val": "single", "sz": 8, "color": LINE}, bottom={"val": "single", "sz": 8, "color": LINE}, left={"val": "single", "sz": 8, "color": LINE}, right={"val": "single", "sz": 8, "color": LINE})
            if row_index + col_index < len(cards):
                title, body = cards[row_index + col_index]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(title)
                r.bold = True
                r.font.size = Pt(10.2)
                r.font.color.rgb = rgb(DEEP_BLUE)
                set_east_asia_font(r)
                p2 = cell.add_paragraph()
                p2.paragraph_format.space_after = Pt(0)
                p2.paragraph_format.line_spacing = 1.15
                r2 = p2.add_run(body)
                r2.font.size = Pt(9.2)
                r2.font.color.rgb = rgb(INK)
                set_east_asia_font(r2)
        prevent_row_split(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_flow(doc):
    labels = [
        ("01", "上传日志", "一个或多个压缩包"),
        ("02", "自动分析", "特征 + 公司 Skill"),
        ("03", "看报告", "证据、建议、导出"),
        ("04", "补充知识", "特征或案例"),
        ("05", "团队复用", "问答、工具、反馈"),
    ]
    table = doc.add_table(rows=1, cols=len(labels))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lock_table_layout(table)
    table.autofit = False
    for idx, (num, title, detail) in enumerate(labels):
        cell = table.cell(0, idx)
        set_column_width(cell, 1.3)
        set_cell_margins(cell, top=95, bottom=95, start=55, end=55)
        set_cell_shading(cell, PALE_BLUE if idx % 2 == 0 else PALE_CYAN)
        set_cell_border(cell, top={"val": "single", "sz": 8, "color": LINE}, bottom={"val": "single", "sz": 8, "color": LINE}, left={"val": "single", "sz": 8, "color": LINE}, right={"val": "single", "sz": 8, "color": LINE})
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(num)
        r1.bold = True
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = rgb(CYAN)
        set_east_asia_font(r1)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(title)
        r2.bold = True
        r2.font.size = Pt(9.8)
        r2.font.color.rgb = rgb(DEEP_BLUE)
        set_east_asia_font(r2)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        p3.paragraph_format.line_spacing = 1.0
        r3 = p3.add_run(detail)
        r3.font.size = Pt(8.2)
        r3.font.color.rgb = rgb(MUTED)
        set_east_asia_font(r3)
    prevent_row_split(table.rows[0])


def add_permissions_table(doc):
    rows = [
        ("贡献者", "上传日志、发布自己的案例/工具、提出需求", "自己的原始日志；公开知识", "不能查看他人的原始日志或复核他人来源"),
        ("审核者", "在贡献者基础上，复核特征来源、下载所需原始资料", "团队范围内依权限开放的资料", "用于质量把关，不代替管理员"),
        ("管理员", "账号、角色和全站治理", "全部受控内容", "仅限被授权的运维人员"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    lock_table_layout(table)
    widths = [0.85, 2.0, 1.55, 2.1]
    headers = ["角色", "你通常可以做什么", "常见可见范围", "注意"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_column_width(cell, widths[i])
        set_cell_shading(cell, DEEP_BLUE)
        set_cell_margins(cell, top=90, bottom=90, start=70, end=70)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(8.7)
        r.font.color.rgb = rgb(WHITE)
        set_east_asia_font(r)
    set_repeat_table_header(table.rows[0])
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            set_column_width(cells[i], widths[i])
            set_cell_shading(cells[i], WHITE if row_index % 2 == 0 else PALE_BLUE)
            set_cell_margins(cells[i], top=80, bottom=80, start=70, end=70)
            set_cell_border(cells[i], top={"val": "single", "sz": 6, "color": LINE}, bottom={"val": "single", "sz": 6, "color": LINE}, left={"val": "single", "sz": 6, "color": LINE}, right={"val": "single", "sz": 6, "color": LINE})
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(value)
            r.font.size = Pt(8.7)
            r.font.color.rgb = rgb(INK)
            if i == 0:
                r.bold = True
                r.font.color.rgb = rgb(DEEP_BLUE)
            set_east_asia_font(r)
        prevent_row_split(table.rows[-1])


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Heiti SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Heiti SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.font.size = Pt(10.2)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_after = Pt(4)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("BMC 智能日志平台  ·  同事使用指导书")
    hr.font.size = Pt(8.3)
    hr.font.color.rgb = rgb(MUTED)
    set_east_asia_font(hr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    fr = fp.add_run("内部使用 · v0.6.0-local · 第 ")
    fr.font.size = Pt(8.1)
    fr.font.color.rgb = rgb(MUTED)
    set_east_asia_font(fr)
    add_page_field(fp)
    fr2 = fp.add_run(" 页")
    fr2.font.size = Pt(8.1)
    fr2.font.color.rgb = rgb(MUTED)
    set_east_asia_font(fr2)


def page_break(doc):
    doc.add_page_break()


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # Cover
    add_label(doc, "Internal quick-start handbook", CYAN)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("BMC 智能日志平台")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = rgb(NAVY)
    set_east_asia_font(r)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(12)
    r2 = p2.add_run("同事使用指导书")
    r2.bold = True
    r2.font.size = Pt(18)
    r2.font.color.rgb = rgb(BLUE)
    set_east_asia_font(r2)
    add_rule(doc, CYAN)
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(12)
    p3.paragraph_format.space_after = Pt(12)
    r3 = p3.add_run("从上传一份日志，到沉淀一条能被团队复用的知识")
    r3.font.size = Pt(12)
    r3.font.color.rgb = rgb(MUTED)
    set_east_asia_font(r3)
    add_image(doc, ASSET_DIR / "guide-cover.png", 6.5, "把重复的日志翻查，变成团队共同的诊断能力。", "工程师在数据中心工作站查看 BMC 诊断界面")
    add_callout(doc, "适合谁读？", "第一次使用平台的硬件维护同事、需要上传问题日志的同事、准备沉淀规则/案例/工具的贡献者。全程不用写代码。", "blue")
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_before = Pt(10)
    p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r4 = p4.add_run("内部使用 · 平台版本 v0.6.0-local · 2026 年 7 月")
    r4.font.size = Pt(8.7)
    r4.font.color.rgb = rgb(MUTED)
    set_east_asia_font(r4)
    page_break(doc)

    # Page 2: three-minute map
    add_label(doc, "First, know the map")
    add_heading(doc, "先用 3 分钟认识这个平台")
    add_para(doc, "你不需要一次学会所有功能。先把它理解成一个“日志分析 + 经验沉淀 + 团队协作”的工作台：上传日志得到报告；把已确认经验变成特征或案例；下次遇到相似问题，平台会更快给出线索。", 10.5)
    add_module_cards(doc, [
        ("日志分析", "上传一个或多个日志压缩包。每个文件会形成一份独立报告，支持常见 tar.gz、tgz、tar、zip、gz 与文本日志。"),
        ("问题报告", "查看命中特征、公司 BMC Skill（如已配置）、证据位置与处置建议；需要时导出独立 HTML 报告。"),
        ("特征库", "把“哪些文件出现什么现象，就可能是什么问题”保存为团队规则，并保留规则版本和历史证据。"),
        ("案例库", "把一次排障经验写成文章，或导入已有 Word、PPT、Markdown、文本资料后再校对发布。"),
        ("智能问答", "用自然语言问定位思路；回答会结合已发布案例和特征，并以流式方式逐步展示。"),
        ("工具中心与需求板", "分享内部小工具与版本说明，也可以提出希望平台新增的能力。"),
    ])
    add_heading(doc, "最推荐的新手路径", 2)
    add_flow(doc)
    add_callout(doc, "你只需记住 4 个词", "原始日志用于追溯；分析报告用于判断；特征用于自动识别；案例用于讲清楚“为什么”和“怎么处理”。", "green")
    page_break(doc)

    # Page 3: account and safety
    add_label(doc, "Before you upload")
    add_heading(doc, "第一步：登录与安全边界")
    add_steps(doc, [
        ("拿到平台地址和账号", "请使用管理员提供的公司内网地址。当前平台默认不开放自助注册；没有账号时，联系管理员开通即可。"),
        ("登录后先看右上角帮助", "每个页面右上角都有“帮助”按钮，第一次进入某个功能时先点一下，就能看到该页应填什么、何时使用。"),
        ("确认你能看到的按钮", "有些按钮只对审核者或管理员显示，例如复核他人的来源日志、调整账号角色。看不到通常是角色权限，而不是操作失败。"),
    ])
    add_heading(doc, "上传前，先做这 3 个判断", 2)
    add_bullets(doc, [
        "确认日志可在公司内网合规留存：不要上传不该进入本团队实例的客户保密资料、个人敏感信息或无关附件。",
        "尽量保留原始压缩包，不要为了“看起来干净”先删目录；文件路径往往正是定位规则的一部分。",
        "写清楚设备/现场背景：型号、BMC 或固件版本、故障现象和发生时间，会显著提高人工复核与问答效率。",
    ])
    add_callout(doc, "重要：分析建议不是最终裁决", "平台会同时给出特征规则、公司 BMC Skill（如已配置）和 LLM 的建议。LLM 只帮助整理思路；请以原始证据、现场状态和已验证规则作为最终判断依据。", "gold")
    add_heading(doc, "角色一眼看懂", 2)
    add_permissions_table(doc)
    page_break(doc)

    # Page 4: batch analysis
    add_label(doc, "Most common task")
    add_heading(doc, "第二步：批量上传 BMC 日志并开始分析")
    add_para(doc, "这是最常用的入口。一次可选择多个压缩包；平台会分别保存、解包分析并生成多份报告。不要等所有日志都完全一样才上传——先让平台给出第一轮线索。", 10.5)
    add_steps(doc, [
        ("进入“日志分析”", "在左侧导航选择“日志分析”，点击上传区域或选择文件。优先选择完整的原始归档，例如 .tar.gz；也支持 .tgz、.tar、.zip、.gz 和单个文本日志。"),
        ("一次选择多个文件", "按住 Ctrl/Command 多选，或直接批量拖入。每份文件会单独生成分析记录，便于分别追溯和下载。"),
        ("补充背景并提交", "如果页面提供描述字段，写上设备标识（脱敏）、现象、时间范围和你最关心的问题。提交后不要关闭页面，等待每份日志状态变为已完成。"),
        ("从列表打开报告", "完成后点击对应报告。平台会自动用当前特征库比对；如部署了公司 BMC Skill，也会并行输出复杂规则的解析结果。"),
    ])
    add_callout(doc, "上传失败时先别反复点", "先检查文件是否仍是完整压缩包、文件是否过大或受网络影响；再打开页面右上角帮助或联系管理员。保留原包能让后续排查更准确。", "red")
    page_break(doc)

    # Page 5 reports
    add_label(doc, "Read the evidence, not only the conclusion")
    add_heading(doc, "第三步：把报告读成可执行的排障线索")
    add_image(doc, ASSET_DIR / "workflow.png", 6.45, "上传不是终点：报告中的证据会继续沉淀为特征、案例和可复用的团队知识。", "压缩日志经上传和检查后形成报告，再连接到规则和案例知识")
    add_heading(doc, "报告页优先看什么？", 2)
    add_steps(doc, [
        ("先看问题概览与严重程度", "把它当作“排查优先级”，不是已确认的故障结论。高严重度先结合现场状态核验。"),
        ("再看证据位置", "重点确认命中的文件路径、原文片段、行号或条件。若证据对不上，请不要直接采纳结论。"),
        ("区分三类来源", "特征库是团队维护的文本/正则规则；BMC Skill 是公司受控的复杂结构化解析（如已配置）；LLM 是辅助解释与定位建议。三者共同参考、互不替代。"),
        ("需要交付时导出 HTML", "报告页可导出独立 HTML 问题分析报告，适合发送给协作同事或放入工单附件。原始归档下载受角色限制，通常仅上传者、审核者和管理员可取。"),
    ])
    add_callout(doc, "一个好习惯", "当你确认“这个结果是对的”，马上补充特征或案例；当你确认“这个结果不对”，记录原因并联系规则维护者。这样下一位同事就不用从头翻日志。", "green")
    page_break(doc)

    # Page 6 features
    add_label(doc, "Turn an observation into a rule")
    add_heading(doc, "第四步：把已确认的问题沉淀到特征库")
    add_para(doc, "特征库的目的不是堆关键词，而是让平台在未来日志中稳定地识别相似现象。一个特征可关联多份日志；一份日志也可对应多个特征。", 10.5)
    add_steps(doc, [
        ("从“特征库”新建", "填写问题名称、严重程度、关键字、描述、诊断和建议。名称尽量使用团队能统一理解的表达，例如“风扇转速异常导致散热风险”。"),
        ("写清楚命中条件", "把条件写成“哪个文件路径 + 出现的关键字或正则 + 条件之间是同时满足还是满足其一”。先写最稳定、最能说明问题的证据。"),
        ("关联来源日志", "上传或选择能证明该问题的日志记录。平台会帮助在关联日志中复核：是某一个文件、多个文件，还是同一份日志中的多个现象共同指向它。"),
        ("阅读 AI 复核建议并人工确认", "AI 可以提示可能的文件与语句，但请逐项检查后再保存。涉及其他人来源日志的复核，需要相应审核权限。"),
        ("保存后观察重新检查结果", "每次保存都会形成新的规则版本，并按当前规则重新检查已完成历史报告。旧命中证据保留，便于回溯。"),
    ])
    add_callout(doc, "特征写作模板", "【名称】现象 + 部件/风险；【条件】明确文件与关键字；【描述】何时出现；【诊断】为什么关联；【建议】下一步检查或更换动作。", "blue")
    page_break(doc)

    # Page 7 cases
    add_label(doc, "Make experience reusable")
    add_heading(doc, "第五步：发布一篇案例，让经验能被搜到")
    add_para(doc, "案例适合保存完整的排障过程：背景、现象、证据、处理过程、验证结果和注意事项。案例正文与特征都会成为智能问答的知识来源。", 10.5)
    add_steps(doc, [
        ("选择“案例库”后新建或导入", "可以直接新建文章，也可以把已经写好的 .docx、.pptx、.md 或 .txt 资料一键导入为草稿。旧版 .doc/.ppt 的导入需要服务器配置 LibreOffice。"),
        ("选一种更舒服的写法", "熟悉排版就选富文本，用工具栏加粗、列表和代码块；习惯技术文档就选 Markdown，可查看语法说明和实时预览。两种方式最后都会统一保存为 Markdown。"),
        ("按结构补齐关键信息", "建议依次写：故障背景 → 现象与影响 → 关键证据 → 排查过程 → 根因/处理 → 验证结果 → 注意事项。可关联多份日志和已有特征。"),
        ("校对草稿再正式发布", "导入的资料先作为待发布草稿，检查标题、图片、敏感信息和结论后再发布。原始导入资料会保留，便于日后追溯。"),
    ])
    add_callout(doc, "案例标题的小技巧", "用“对象 + 现象 + 结果”命名，例如“R220KV2 风扇告警：从 SEL 事件确认单风扇失效”。标题越具体，问答检索和同事搜索越容易命中。", "green")
    page_break(doc)

    # Page 8 QA and tools
    add_label(doc, "Ask, share, improve")
    add_heading(doc, "第六步：用智能问答、工具中心和需求板协作")
    add_heading(doc, "智能问答：把问题问具体", 2)
    add_bullets(doc, [
        "推荐问法：" + "“某型号服务器出现风扇相关告警，报告里已命中特征 A，接下来应优先核对哪些文件和现场项？”",
        "尽量带上：设备/版本、现象、报告中的证据或错误码、你已经做过的操作、希望得到的输出（排查顺序、可能根因、案例链接）。",
        "回答会逐步流式出现并渲染 Markdown；历史会按当前账号保存，适合回顾上次讨论。不要把模型输出当成未验证的维修指令。",
    ])
    add_heading(doc, "工具中心：让小工具不再“只在一个人的电脑里”", 2)
    add_bullets(doc, [
        "发布工具时写清工具版本、用途、作者、使用条件、风险提示和最小示例；需要时上传 Logo。",
        "新版本应作为更新留存，历史上传记录与下载计数不会被清零。下载同事可在工具详情留下改进建议。",
    ])
    add_heading(doc, "需求留言板：把“希望有个功能”写成可评估需求", 2)
    add_bullets(doc, [
        "说清楚“谁在什么场景遇到什么麻烦”，再写希望的结果。例如：批量分析后，希望能按严重度导出一张汇总表。",
        "尽量附上脱敏样例或流程；团队可按周度或半月度筛选价值高、可落地的需求。",
    ])
    add_callout(doc, "保护团队知识", "不要在案例、问答或需求里粘贴不应共享的密钥、账号口令、客户敏感信息或完整内网规则。对真实生产规则请遵循公司的受控流程。", "red")
    page_break(doc)

    # Page 9 scenarios
    add_label(doc, "Copy these scenarios")
    add_heading(doc, "三个常见场景：照着做就可以")
    add_heading(doc, "场景 A｜我刚拿到一份 BMC 日志，想知道有没有明显问题", 2)
    add_steps(doc, [
        ("上传原始归档", "进入日志分析，选择完整 .tar.gz 或其他支持格式；如果有多台机器，批量选择即可。"),
        ("打开每份完成报告", "先按严重程度看优先级，再逐条打开证据位置，确认是否符合现场现象。"),
        ("导出或继续追问", "需要转交时导出 HTML；不确定时到智能问答带着证据问“下一步应该检查什么”。"),
    ])
    add_heading(doc, "场景 B｜我已经确认一个新问题，想让下次自动提示", 2)
    add_steps(doc, [
        ("保留能说明问题的日志", "不要只截图，关联原始报告。确认关键文件、关键字、时间关系和误报边界。"),
        ("在特征库创建条件", "写清问题名、严重程度、诊断、建议和多条匹配条件。让 AI 提示候选证据，但由你最终确认。"),
        ("保存后检查历史报告", "查看重新检查是否符合预期；若误报，收紧路径、关键字或组合条件，再保存新版本。"),
    ])
    add_heading(doc, "场景 C｜我有一篇 Word 排障记录，想让团队以后能搜到", 2)
    add_steps(doc, [
        ("导入为草稿", "进入案例库，导入 Word/PPT/Markdown/文本资料。"),
        ("整理为统一结构", "补全背景、证据、根因、处理、验证和注意事项；删除不应长期留存的敏感内容。"),
        ("关联特征/日志后发布", "让问答和特征库都能发现它；之后同事可读文章、下载来源资料（按权限）并复用经验。"),
    ])
    page_break(doc)

    # Page 10 troubleshooting
    add_label(doc, "When something does not look right")
    add_heading(doc, "常见问题与最快处理方式")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    lock_table_layout(table)
    widths = [2.0, 2.25, 2.25]
    for i, header in enumerate(["你遇到的情况", "先自己检查", "仍无法解决时"]):
        cell = table.cell(0, i)
        set_column_width(cell, widths[i])
        set_cell_shading(cell, DEEP_BLUE)
        set_cell_margins(cell, top=90, bottom=90, start=80, end=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = rgb(WHITE)
        set_east_asia_font(r)
    set_repeat_table_header(table.rows[0])
    troubles = [
        ("登录后看不到上传/复核按钮", "确认所在页面与账号角色；某些操作只对审核者或管理员开放。", "截取页面名称和按钮区域，联系管理员核对角色。"),
        ("日志一直没有完成分析", "确认网络、归档是否完整、文件是否包含可读日志；等待批量任务依次完成。", "提供日志文件名、上传时间和页面状态，不要反复重复上传。"),
        ("报告结论和现场不一致", "回到证据位置，核对文件、行号、版本和时间；区分规则、Skill、LLM 三类来源。", "把不一致点记录为案例/需求或提交给规则维护者，避免直接按结论执行。"),
        ("问答没有给出想要的答案", "补充设备型号、错误码、报告证据和期望输出；问题越具体越好。", "检查是否有相关案例/特征；必要时新建案例补充团队知识。"),
        ("导入的 Word/PPT 排版不理想", "回到草稿，按背景—证据—处理—验证重新整理；检查图片和敏感信息。", "保留原资料并说明格式问题；旧版 .doc/.ppt 需管理员确认服务器转换环境。"),
    ]
    for idx, row_data in enumerate(troubles):
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_column_width(cell, widths[i])
            set_cell_shading(cell, WHITE if idx % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell, top=85, bottom=85, start=80, end=80)
            set_cell_border(cell, top={"val": "single", "sz": 6, "color": LINE}, bottom={"val": "single", "sz": 6, "color": LINE}, left={"val": "single", "sz": 6, "color": LINE}, right={"val": "single", "sz": 6, "color": LINE})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(value)
            r.font.size = Pt(8.9)
            r.font.color.rgb = rgb(INK)
            if i == 0:
                r.bold = True
                r.font.color.rgb = rgb(DEEP_BLUE)
            set_east_asia_font(r)
        prevent_row_split(row)
    add_callout(doc, "求助时请一次带齐这 5 项", "页面名称、账号角色、操作时间、日志/报告编号或文件名、你看到的提示或截图。这样管理员或规则维护者能快速复现，不用来回猜。", "blue")
    page_break(doc)

    # Page 11 checklist / close
    add_label(doc, "One-page checklist")
    add_heading(doc, "最后一页：每次使用前后的检查清单")
    add_heading(doc, "上传前", 2)
    add_bullets(doc, [
        "我上传的是完整、合规留存的原始日志，且没有混入不应共享的敏感资料。",
        "我知道这台设备的基本背景：型号/版本/时间范围/现场现象。",
        "如果是批量上传，我能在每份日志完成后分别核对报告。",
    ])
    add_heading(doc, "看报告时", 2)
    add_bullets(doc, [
        "我不仅看了结论，还核对了文件路径、原文证据和行号/条件。",
        "我分清了特征库、公司 BMC Skill 与 LLM 建议的来源。",
        "我把已确认结果导出 HTML 或写入工单，便于协作与追溯。",
    ])
    add_heading(doc, "确认问题后", 2)
    add_bullets(doc, [
        "可自动识别的模式：我会创建或改进特征，并关联来源日志。",
        "值得复用的排障过程：我会发布案例，写清证据、处理和验证。",
        "平台本身缺少的能力：我会在需求留言板中描述真实场景与预期结果。",
    ])
    add_callout(doc, "平台越用越聪明，靠的是每一次“确认后沉淀”", "上传日志解决的是今天的问题；维护特征和案例，解决的是未来团队重复遇到的问题。感谢你把经验留给下一位同事。", "green")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("需要帮助？请先使用页面右上角“帮助”，再联系平台管理员。")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = rgb(DEEP_BLUE)
    set_east_asia_font(r)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("帮助署名：作者：江小伟 00983211")
    r2.font.size = Pt(9)
    r2.font.color.rgb = rgb(MUTED)
    set_east_asia_font(r2)

    doc.core_properties.title = "BMC 智能日志平台同事使用指导书"
    doc.core_properties.subject = "面向服务器硬件维护团队的新人上手手册"
    doc.core_properties.author = "江小伟 00983211"
    doc.core_properties.comments = "Generated from docs/build_user_guide.py"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
